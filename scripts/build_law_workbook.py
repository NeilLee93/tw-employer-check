"""產出「法條白話化與嚴重度對照表」的填寫用 Excel（待辦第 4 項）。

這支腳本不做判讀，只把實際出現過的法條與官方描述整理成可填的表格。
白話化與嚴重度需要勞動法專業，由人填，不由程式猜。

執行：
    python scripts/build_law_workbook.py            # 產出 data/勞基法白話對照表.xlsx
    python scripts/build_law_workbook.py --force    # 覆寫已存在的檔案（會蓋掉手填欄位）

預設**拒絕覆寫**已存在的檔案——這份表填過之後就是專案最貴的資產。
"""

from __future__ import annotations

import collections
import csv
import re
import statistics
import sys
import unicodedata
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data" / "lsa_violations_raw.csv"
OUT = ROOT / "data" / "勞基法白話對照表.xlsx"

# 前 8 條吃掉 87% 的引用，本輪只做這 8 條。順序即引用量。
TOP_ARTICLES = [24, 30, 32, 36, 22, 39, 23, 38]
ARTICLE_TOPIC = {
    24: "延長工時工資（加班費）",
    30: "正常工時與出勤紀錄",
    32: "延長工時上限與程序",
    36: "例假與休息日",
    22: "工資給付方式",
    39: "假日工資",
    23: "工資給付與工資清冊",
    38: "特別休假",
}

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
FILL_FILL = PatternFill("solid", fgColor="FFF2CC")  # 要人填的欄位染色
HEADER_FONT = Font(color="FFFFFF", bold=True)


def clean(s: str | None) -> str:
    """全形轉半形、去掉所有空白與換行。欄位裡兩者都有。"""
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", s or ""))


def clean_text(s: str | None) -> str:
    """官方描述另外去掉句尾標點。

    「延長工作時間未依規定加給工資」與同一句加上句號，是同一種違規的兩種寫法，
    不合併會讓要改寫的句子平白多出上百句。
    """
    return clean(s).rstrip("。.、,;；:：")


def normalize_law(raw: str) -> str:
    """把法條字串收斂成 `勞動基準法第X條第Y項`。

    兩個坑：「勞基法」與「勞動基準法」混用；少數列尾巴重複串接法規名稱
    （`勞動基準法第24條勞動基準法勞動基準法`，全量 62 對）。
    """
    s = clean(raw).replace("勞基法", "勞動基準法")
    s = re.sub(r"(勞動基準法)+", "勞動基準法", s)
    if s != "勞動基準法":
        s = re.sub(r"勞動基準法$", "", s)
    return s


def article_of(law: str) -> int | None:
    m = re.match(r"^勞動基準法第(\d+)條", law)
    return int(m.group(1)) if m else None


def load():
    """展開成（法條, 官方描述, 罰鍰, 是否單條列）。

    `違法法規法條` 與 `違反法規內容` 是兩串平行的 `;` 分隔字串，
    數量對不上的列（3.87%）無法確定哪句對哪條，整列跳過不猜。
    """
    pairs = []
    mismatched = 0
    total_rows = 0
    with open(SOURCE, encoding="utf-8", errors="replace", newline="") as fh:
        for row in csv.DictReader(fh):
            total_rows += 1
            laws = (row.get("違法法規法條") or "").split(";")
            texts = (row.get("違反法規內容") or "").split(";")
            if len(laws) != len(texts):
                mismatched += 1
                continue
            fine = (row.get("罰鍰金額") or "").strip()
            fine_value = int(fine) if fine.isdigit() else None
            for law, text in zip(laws, texts):
                law = normalize_law(law)
                if law:
                    pairs.append((law, clean_text(text), fine_value, len(laws) == 1))
    return pairs, total_rows, mismatched


def style_header(ws, widths, fill_from=None) -> None:
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    if fill_from:
        start = ws[f"{fill_from}1"].column
        for row in ws.iter_rows(min_row=2, min_col=start, max_col=ws.max_column):
            for cell in row:
                cell.fill = FILL_FILL


def sheet_readme(wb, stats) -> None:
    ws = wb.create_sheet("說明")
    ws.column_dimensions["A"].width = 110
    lines = [
        ("怎麼填這份表", "title"),
        ("", ""),
        ("這份表是待辦第 4 項的輸入。黃色欄位要人填，其餘是從資料算出來的，不要改。", ""),
        ("", ""),
        ("填寫順序", "h2"),
        (f"1. 先填「條項層」——{stats['clauses']} 列，是必填。填完就能做出第一版判讀。", ""),
        (f"2. 有餘力再填「句型層」——{stats['phrases']:,} 種句型裡，"
         f"標「必」的 {stats['must']} 句涵蓋各條前 80%，"
         f"再加標「宜」的 {stats['good']} 句就到 90%。", ""),
        ("   剩下的長尾多是同義改寫與個案細節，可以不理。", ""),
        ("3. 「其餘條文」是本輪不做的 30 條，供你決定第二輪要不要往下做。", ""),
        ("", ""),
        ("嚴重度尺度（草案，請直接改）", "h2"),
        ("5　勞工的錢或命當場受損且難以回復：職災、工資完全未給付", ""),
        ("4　法定給付短少：加班費、假日工資、特休未給或短給", ""),
        ("3　超時與休息被侵蝕：延長工時超過上限、例假未給", ""),
        ("2　程序與制度瑕疵：未經工會或勞資會議同意、未公告", ""),
        ("1　形式義務未落實：出勤紀錄、工資清冊記載不全", ""),
        ("", ""),
        ("資料範圍與已知限制", "h2"),
        (f"資料：勞動部「違反勞動法令事業單位（勞動基準法）」，{stats['total_rows']:,} 列。", ""),
        (f"展開後的（法條, 官方描述）對共 {stats['pairs']:,} 個。", ""),
        (f"法條與描述數量對不上、無法確定哪句對哪條的列有 {stats['mismatched']:,} 列"
         f"（{stats['mismatched'] / stats['total_rows']:.2%}），已整列跳過，不猜。", ""),
        (f"本輪的 8 條涵蓋勞基法全部引用的 {stats['top_share']:.1%}。", ""),
        ("罰鍰欄位 2020 年以前全空，且一列可能同時違反多條。表上的罰鍰統計"
         "只取「整列只引用這一條」的列，才能確定罰鍰是這一條造成的。", ""),
        ("同一條的不同「項」分開列，因為嚴重度不同（例如第 30 條第 5、6 項是"
         "出勤紀錄的記載與保存，第 1 項才是工時上限本身）。", ""),
        ("", ""),
        ("⚠ 重跑 scripts/build_law_workbook.py 會覆寫本檔，手填欄位會被蓋掉。"
         "腳本預設拒絕覆寫，要覆寫得自己加 --force。", "warn"),
    ]
    for i, (text, kind) in enumerate(lines, start=1):
        cell = ws.cell(row=i, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if kind == "title":
            cell.font = Font(size=14, bold=True)
        elif kind == "h2":
            cell.font = Font(size=11, bold=True, color="1F3864")
        elif kind == "warn":
            cell.font = Font(bold=True, color="C00000")


def sheet_clauses(wb, pairs) -> None:
    """條項層：本輪必填的 30 列。"""
    ws = wb.create_sheet("條項層（必填）")
    ws.append([
        "條", "主題", "法條", "引用次數", "佔勞基法全體",
        "句型數", "單條列數", "罰鍰中位數", "罰鍰平均",
        "最常見的官方描述", "白話說明（一句話）", "嚴重度", "判讀備註",
    ])

    total = len(pairs)
    by_clause = collections.defaultdict(list)
    for law, text, fine, alone in pairs:
        by_clause[law].append((text, fine, alone))

    rows = [
        (law, items)
        for law, items in by_clause.items()
        if article_of(law) in TOP_ARTICLES
    ]
    rows.sort(key=lambda kv: (TOP_ARTICLES.index(article_of(kv[0])), -len(kv[1])))

    for law, items in rows:
        article = article_of(law)
        fines = [f for _, f, alone in items if alone and f]
        texts = collections.Counter(t for t, _, _ in items)
        ws.append([
            f"第 {article} 條",
            ARTICLE_TOPIC[article],
            law,
            len(items),
            len(items) / total,
            len(texts),
            sum(1 for _, _, alone in items if alone),
            statistics.median(fines) if fines else None,
            round(statistics.mean(fines)) if fines else None,
            texts.most_common(1)[0][0],
            None, None, None,
        ])

    for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        row[0].number_format = "0.0%"
    for row in ws.iter_rows(min_row=2, min_col=8, max_col=9):
        for cell in row:
            cell.number_format = "#,##0"

    style_header(ws, {
        "A": 9, "B": 22, "C": 26, "D": 10, "E": 13, "F": 8, "G": 10,
        "H": 12, "I": 12, "J": 34, "K": 42, "L": 8, "M": 30,
    }, fill_from="K")
    add_severity_validation(ws, "L")


def sheet_phrases(wb, pairs) -> None:
    """句型層：官方描述去重之後的實際句子。白話化是改寫這些句子。"""
    ws = wb.create_sheet("句型層（選填）")
    ws.append([
        "條", "法條", "官方描述（改寫這一欄）", "次數",
        "佔該條", "該條累積", "建議填寫", "白話改寫", "嚴重度微調", "備註",
    ])

    counter = collections.Counter()
    for law, text, _, _ in pairs:
        article = article_of(law)
        if article in TOP_ARTICLES:
            counter[(article, law, text)] += 1

    by_article = collections.defaultdict(list)
    for (article, law, text), n in counter.items():
        by_article[article].append((n, law, text))

    for article in TOP_ARTICLES:
        items = sorted(by_article[article], key=lambda x: (-x[0], x[2]))
        total = sum(n for n, _, _ in items)
        cumulative = 0
        for n, law, text in items:
            cumulative += n
            share = cumulative / total
            hint = "必" if share <= 0.8 else ("宜" if share <= 0.9 else "")
            ws.append([
                f"第 {article} 條", law, text, n,
                n / total, share, hint, None, None, None,
            ])

    for row in ws.iter_rows(min_row=2, min_col=5, max_col=6):
        for cell in row:
            cell.number_format = "0.0%"

    style_header(ws, {
        "A": 9, "B": 26, "C": 52, "D": 8, "E": 9, "F": 10,
        "G": 10, "H": 46, "I": 11, "J": 24,
    }, fill_from="H")
    add_severity_validation(ws, "I")


def sheet_rest(wb, pairs) -> None:
    """本輪不做的其餘條文，供第二輪選題。"""
    ws = wb.create_sheet("其餘條文（參考）")
    ws.append(["法條", "引用次數", "佔勞基法全體", "句型數", "最常見的官方描述"])

    total = len(pairs)
    by_article = collections.defaultdict(list)
    for law, text, _, _ in pairs:
        article = article_of(law)
        if article is not None and article not in TOP_ARTICLES:
            by_article[article].append(text)

    for article, texts in sorted(by_article.items(), key=lambda kv: -len(kv[1])):
        counter = collections.Counter(texts)
        ws.append([
            f"勞動基準法第 {article} 條",
            len(texts),
            len(texts) / total,
            len(counter),
            counter.most_common(1)[0][0],
        ])
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].number_format = "0.0%"
    style_header(ws, {"A": 24, "B": 11, "C": 14, "D": 8, "E": 60})


def add_severity_validation(ws, column: str) -> None:
    """嚴重度限制在 1-5，避免填出 3.5 或「高」這種下游算不動的值。"""
    dv = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True)
    dv.error = "嚴重度請填 1 到 5 的整數"
    ws.add_data_validation(dv)
    dv.add(f"{column}2:{column}{ws.max_row}")


def main() -> None:
    if OUT.exists() and "--force" not in sys.argv:
        sys.exit(
            f"{OUT} 已存在。重跑會蓋掉手填欄位；確定要覆寫請加 --force，"
            f"或先把現有檔案改名備份。"
        )
    if not SOURCE.exists():
        sys.exit(f"找不到 {SOURCE}，請先依 README 下載原始資料。")

    pairs, total_rows, mismatched = load()
    top = [p for p in pairs if article_of(p[0]) in TOP_ARTICLES]
    stats = {
        "total_rows": total_rows,
        "pairs": len(pairs),
        "mismatched": mismatched,
        "top_share": len(top) / len(pairs),
    }

    wb = Workbook()
    wb.remove(wb.active)
    sheet_clauses(wb, pairs)
    sheet_phrases(wb, pairs)
    sheet_rest(wb, pairs)

    phrases = wb["句型層（選填）"]
    hints = [r[0].value for r in phrases.iter_rows(min_row=2, min_col=7, max_col=7)]
    stats["clauses"] = wb["條項層（必填）"].max_row - 1
    stats["phrases"] = phrases.max_row - 1
    stats["must"] = hints.count("必")
    stats["good"] = hints.count("宜")

    sheet_readme(wb, stats)
    wb.move_sheet("說明", offset=-3)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)

    must, good = stats["must"], stats["good"]
    print(f"寫出 {OUT}")
    print(f"  條項層（必填）    {wb['條項層（必填）'].max_row - 1:>5} 列")
    print(f"  句型層（選填）    {phrases.max_row - 1:>5} 列（標「必」{must}、標「宜」{good}）")
    print(f"  其餘條文（參考）  {wb['其餘條文（參考）'].max_row - 1:>5} 列")
    print(f"  本輪 8 條涵蓋率   {stats['top_share']:.1%}")


if __name__ == "__main__":
    main()
