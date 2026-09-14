"""判讀層：把 `twec.roster` 歸戶出的一個雇主，套上白話與嚴重度。

輸入是 `Entity`（一個雇主底下所有別名）、違規原始 CSV、與人工填的
「法條白話對照表」（`scripts/build_law_workbook.py` 產出，`data/勞基法白話
對照表.xlsx`）。輸出是一份 `InterpretedReport`：每筆違規套上白話說明與
嚴重度，加上累犯與行政救濟中的彙總。

兩條 fallback 規則是待辦第 4 項卡住之後、拿 spike/06_interpret_logic.py
在真實資料上驗證過的（見該檔開頭的完整推導）：

    白話：人工填了就用人工的；沒填就用**這一列自己的官方描述文字**，
    不是表格裡「這條底下最常見」的那句——那句不保證跟這一列實際發生的
    事對得上。

    嚴重度：人工填了就用人工的；沒填就用罰鍰**平均**（不是中位數）在
    「有罰鍰平均資料的條項」裡切五級分位數。中位數在這份資料上鑑別力
    不夠：40 條項裡 27 條中位數都卡在法定最低罰鍰 20,000（多數個案沒有
    加重情節），quantiles 幾乎全部撞在同一格。平均會被累犯/加重情節的
    個案拖高，才是嚴重度真正想抓的信號。完全沒罰鍰資料的條項標
    「未評級（無罰鍰資料）」；法條根本不在表格裡的標
    「未評級（不在表格範圍）」——不假裝有判讀。

歸戶：吃 `Entity.names`（該雇主所有別名，含統編併起來的更名組），
一次掃過 CSV 用集合比對，不是對每個別名各掃一次全檔。
"""

from __future__ import annotations

import bisect
import csv
import os
import re
import statistics
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from twec.names import normalize as normalize_name
from twec.names import split_entity
from twec.roster import Entity

NAME_COLUMN = "事業單位名稱或負責人"
DISPUTED_MARKERS = ("行政救濟", "訴願", "撤銷")


def _clean(s: str | None) -> str:
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", s or ""))


def _clean_text(s: str | None) -> str:
    return _clean(s).rstrip("。.、,;；:：")


def normalize_law(raw: str) -> str:
    """把法條字串收斂成 `勞動基準法第X條第Y項`。

    兩個坑：「勞基法」與「勞動基準法」混用；少數列尾巴重複串接法規名稱
    （`勞動基準法第24條勞動基準法勞動基準法`，全量 62 對）。
    """
    s = _clean(raw).replace("勞基法", "勞動基準法")
    s = re.sub(r"(勞動基準法)+", "勞動基準法", s)
    if s != "勞動基準法":
        s = re.sub(r"勞動基準法$", "", s)
    return s


def article_of(law: str) -> int | None:
    m = re.match(r"^勞動基準法第(\d+)條", law)
    return int(m.group(1)) if m else None


@dataclass(frozen=True)
class LawEntry:
    """法條白話對照表「條項層」的一列，用 `法條` 字串當 key。"""

    law: str
    topic: str
    plain: str | None
    manual_severity: int | None
    fine_median: float | None
    fine_mean: float | None
    official_text: str


@dataclass(frozen=True)
class InterpretedItem:
    """一列違規裁處，套用白話與嚴重度之後的樣子。"""

    date: str
    agency: str
    case_no: str
    law: str
    article: int | None
    text: str
    text_source: str
    severity: int | None
    severity_source: str
    fine: int | None
    disputed: bool
    note: str


@dataclass
class InterpretedReport:
    entity: Entity
    items: list[InterpretedItem]
    repeat_offenses: dict[str, int] = field(default_factory=dict)
    disputed_count: int = 0

    @property
    def top_severity(self) -> int | None:
        graded = [i.severity for i in self.items if i.severity is not None]
        return max(graded) if graded else None

    @property
    def distinct_laws(self) -> int:
        return len({i.law for i in self.items})


def load_law_table(xlsx_path: str | os.PathLike[str]) -> dict[str, LawEntry]:
    """讀「條項層（必填）」分頁，回傳 法條 -> LawEntry。"""
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb["條項層（必填）"]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    idx = {name: i for i, name in enumerate(header)}
    table: dict[str, LawEntry] = {}
    for r in rows[1:]:
        law = r[idx["法條"]]
        if not law:
            continue
        table[law] = LawEntry(
            law=law,
            topic=r[idx["主題"]] or "",
            plain=r[idx["白話說明（一句話）"]] or None,
            manual_severity=r[idx["嚴重度"]] or None,
            fine_median=r[idx["罰鍰中位數"]],
            fine_mean=r[idx["罰鍰平均"]],
            official_text=r[idx["最常見的官方描述"]] or "",
        )
    return table


def build_severity_buckets(law_table: dict[str, LawEntry]) -> list[float]:
    """用表格裡「有罰鍰平均」的條項，切出五級分位數邊界（4 條分界線）。

    見模組開頭：中位數在這份資料上幾乎全部卡在法定最低罰鍰 20,000，
    沒有鑑別力，改用平均。
    """
    means = sorted(e.fine_mean for e in law_table.values() if e.fine_mean)
    if len(means) < 5:
        return []
    return statistics.quantiles(means, n=5)


def severity_from_fine(fine_mean: float, buckets: list[float]) -> int:
    """罰鍰平均落在分位數第幾段，就是嚴重度 1-5（段數低＝罰得輕＝分數低）。"""
    return bisect.bisect_right(buckets, fine_mean) + 1


def resolve_text(entry: LawEntry | None, row_text: str) -> tuple[str, str]:
    """白話 fallback：人工白話 → 這一列自己的官方描述文字。"""
    if entry is not None and entry.plain:
        return entry.plain, "人工白話"
    return row_text, "官方描述（原始文字，尚無白話）"


def resolve_severity(entry: LawEntry | None, buckets: list[float]) -> tuple[int | None, str]:
    if entry is None:
        return None, "未評級（不在表格範圍）"
    if entry.manual_severity is not None:
        return entry.manual_severity, "人工"
    if entry.fine_mean and buckets:
        return severity_from_fine(entry.fine_mean, buckets), "罰鍰平均推算"
    return None, "未評級（無罰鍰資料）"


def parse_violation_row(row: dict[str, str]) -> list[tuple[str, str, int | None]] | None:
    """一列違規 CSV 展開成 [(法條, 官方描述, 罰鍰)]。數量對不上回傳 None（整列跳過，不猜）。"""
    laws = (row.get("違法法規法條") or "").split(";")
    texts = (row.get("違反法規內容") or "").split(";")
    if len(laws) != len(texts):
        return None
    fine = (row.get("罰鍰金額") or "").strip()
    fine_value = int(fine) if fine.isdigit() else None
    out = []
    for law, text in zip(laws, texts):
        law = normalize_law(law)
        if law:
            out.append((law, _clean_text(text), fine_value))
    return out


def interpret_entity(
    entity: Entity,
    raw_csv_path: str | os.PathLike[str],
    law_table: dict[str, LawEntry],
    buckets: list[float],
) -> InterpretedReport:
    """查一個雇主（含所有別名）在違規名單裡的所有列，套用白話與嚴重度。"""
    aliases = set(entity.names)
    items: list[InterpretedItem] = []
    law_counter: dict[str, int] = {}
    disputed = 0

    with open(raw_csv_path, encoding="utf-8", errors="replace", newline="") as fh:
        for row in csv.DictReader(fh):
            raw_name = row.get(NAME_COLUMN) or ""
            if not raw_name.strip():
                continue
            if split_entity(normalize_name(raw_name)).org not in aliases:
                continue

            parsed = parse_violation_row(row)
            if parsed is None:
                continue

            note = (row.get("備註說明") or "").strip()
            is_disputed = any(m in note for m in DISPUTED_MARKERS)
            if is_disputed:
                disputed += 1

            for law, text, fine in parsed:
                law_entry = law_table.get(law)
                display_text, text_source = resolve_text(law_entry, text)
                severity, severity_source = resolve_severity(law_entry, buckets)
                law_counter[law] = law_counter.get(law, 0) + 1
                items.append(
                    InterpretedItem(
                        date=row.get("處分日期") or "",
                        agency=row.get("主管機關") or "",
                        case_no=row.get("處分字號") or "",
                        law=law,
                        article=article_of(law),
                        text=display_text,
                        text_source=text_source,
                        severity=severity,
                        severity_source=severity_source,
                        fine=fine,
                        disputed=is_disputed,
                        note=note,
                    )
                )

    items.sort(key=lambda i: i.date, reverse=True)
    repeat = {law: n for law, n in law_counter.items() if n > 1}
    return InterpretedReport(
        entity=entity,
        items=items,
        repeat_offenses=repeat,
        disputed_count=disputed,
    )
